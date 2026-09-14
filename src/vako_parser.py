import requests
from lxml import html
import openpyxl
from openpyxl.styles import Font
from datetime import datetime
import time

# ==========================================================
# 1. КОНФИГУРАЦИЯ — 5 МОДЕЛЕЙ
# ==========================================================

vako_products = [
    {
        "name": "iPhone 17 Pro Max 256GB",
        "url": "https://vako.market/apple-iphone-17-pro-max-256gb",
        "xpath_model": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__title')]//a",
        "xpath_price": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-price')]",
        "xpath_seller": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-detail')]//div[contains(@class, 'minicard__store-name')]//a",
        "xpath_link": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__more')]//div[contains(@class, 'minicard__in-store')]//a"
    },
    {
        "name": "iPhone 17 Pro 256GB",
        "url": "https://vako.market/apple-iphone-17-pro-256gb",
        "xpath_model": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__title')]//a",
        "xpath_price": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-price')]",
        "xpath_seller": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-detail')]//div[contains(@class, 'minicard__store-name')]//a",
        "xpath_link": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__more')]//div[contains(@class, 'minicard__in-store')]//a"
    },
    {
        "name": "iPhone 17 256GB",
        "url": "https://vako.market/apple-iphone-17-256gb",
        "xpath_model": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__title')]//a",
        "xpath_price": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-price')]",
        "xpath_seller": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-detail')]//div[contains(@class, 'minicard__store-name')]//a",
        "xpath_link": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__more')]//div[contains(@class, 'minicard__in-store')]//a"
    },
    {
        "name": "iPhone 16 Pro Max 256GB",
        "url": "https://vako.market/apple-iphone-16-pro-max-256gb",
        "xpath_model": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__title')]//a",
        "xpath_price": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-price')]",
        "xpath_seller": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-detail')]//div[contains(@class, 'minicard__store-name')]//a",
        "xpath_link": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__more')]//div[contains(@class, 'minicard__in-store')]//a"
    },
    {
        "name": "iPhone 17 Pro Max 512GB",
        "url": "https://vako.market/apple-iphone-17-pro-max-512gb",
        "xpath_model": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__title')]//a",
        "xpath_price": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-price')]",
        "xpath_seller": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__store-detail')]//div[contains(@class, 'minicard__store-name')]//a",
        "xpath_link": "//div[contains(@class, 'catalog-category__main')]//div[contains(@class, 'minicard__more')]//div[contains(@class, 'minicard__in-store')]//a"
    }
]

# ==========================================================
# 2. СЛОВАРЬ ПРИОРИТЕТОВ ПО ЦВЕТУ
# ==========================================================

COLOR_PRIORITY = {
    "iPhone 17 Pro Max 256GB": {"Оранжевый": 1, "Синий": 2, "Серебристый": 3},
    "iPhone 17 Pro 256GB": {"Оранжевый": 1, "Синий": 2, "Серебристый": 3},
    "iPhone 17 256GB": {"Чёрный": 1, "Голубой": 2, "Зелёный": 3},
    "iPhone 16 Pro Max 256GB": {"Пустынный титан": 1, "Чёрный титан": 2, "Белый титан": 3},
    "iPhone 17 Pro Max 512GB": {"Оранжевый": 1, "Синий": 2, "Серебристый": 3}
}

# ==========================================================
# 3. ФУНКЦИЯ ИЗВЛЕЧЕНИЯ ЦВЕТА
# ==========================================================

def extract_color(full_name):
    """
    Извлекает цвет из полного названия модели.
    Порядок важен: сначала ищем более конкретные цвета.
    """
    colors = [
        # Титан-варианты (для 16 Pro Max и Pro)
        ("Natural Titanium", "Титан"),
        ("Desert Titanium", "Пустынный титан"),
        ("White Titanium", "Белый титан"),
        ("Black Titanium", "Чёрный титан"),
        ("Titanium", "Титан"),
        ("Пустынный титан", "Пустынный титан"),
        ("Белый титан", "Белый титан"),
        ("Чёрный титан", "Чёрный титан"),
        ("Титан", "Титан"),

        # Цвета 17-й серии
        ("Cosmic Orange", "Оранжевый"),
        ("Deep Blue", "Синий"),
        ("Mist Blue", "Голубой"),
        ("Sage", "Зелёный"),
        ("Silver", "Серебристый"),
        ("Серебристый", "Серебристый"),
        ("Оранжевый", "Оранжевый"),
        ("Синий", "Синий"),
        ("Голубой", "Голубой"),
        ("Зелёный", "Зелёный"),

        # Универсальные цвета
        ("Чёрный", "Чёрный"),
        ("Черный", "Чёрный"),
        ("Black", "Чёрный"),
        ("White", "Белый"),
        ("Белый", "Белый"),
        ("Gold", "Золотой"),
        ("Золотой", "Золотой"),
        ("Pink", "Розовый"),
        ("Розовый", "Розовый"),
        ("Purple", "Фиолетовый"),
        ("Фиолетовый", "Фиолетовый"),
    ]

    name_lower = full_name.lower()
    for keyword, color_name in colors:
        if keyword.lower() in name_lower:
            return color_name
    return "Не указан"

# ==========================================================
# 4. ФУНКЦИЯ ОПРЕДЕЛЕНИЯ ПРИОРИТЕТА ЦВЕТА
# ==========================================================

def get_color_priority(model_name, color):
    priorities = COLOR_PRIORITY.get(model_name, {})
    return priorities.get(color, "—")

# ==========================================================
# 5. ФУНКЦИЯ ПАРСИНГА VAKO.MARKET С RETRY
# ==========================================================

def parse_vako_prices(url, xpath_model, xpath_price, xpath_seller, xpath_link, timeout=15, max_retries=3):
    """
    Парсит все цены с Vako.Market с повторными попытками при ошибках сети.
    """
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
        'Accept-Language': 'ru-RU,ru;q=0.8,en-US;q=0.5,en;q=0.3',
        'Accept-Encoding': 'gzip, deflate, br',
        'Connection': 'keep-alive',
        'Upgrade-Insecure-Requests': '1',
        'Sec-Fetch-Dest': 'document',
        'Sec-Fetch-Mode': 'navigate',
        'Sec-Fetch-Site': 'none',
        'Sec-Fetch-User': '?1',
        'Cache-Control': 'max-age=0',
    }

    for attempt in range(max_retries):
        try:
            session = requests.Session()
            session.headers.update(headers)

            wait_time = 2 + attempt * 3
            time.sleep(wait_time)

            response = session.get(url, timeout=timeout)

            if "Are you human" in response.text or "Security check" in response.text:
                with open("vako_captcha.html", "w", encoding="utf-8") as f:
                    f.write(response.text)
                return [{"error": "Капча (Are you human)"}]

            if response.status_code != 200:
                if attempt < max_retries - 1:
                    continue
                return [{"error": f"HTTP {response.status_code}"}]

            tree = html.fromstring(response.text)

            models = tree.xpath(xpath_model)
            sellers = tree.xpath(xpath_seller)
            prices = tree.xpath(xpath_price)
            links = tree.xpath(xpath_link)

            if not prices or not sellers:
                return [{"error": "Не найдено"}]

            results = []
            count = min(len(models) if models else len(prices), len(sellers), len(prices))

            for i in range(count):
                model_name = ""
                if models and i < len(models):
                    model_el = models[i]
                    model_name = model_el.text.strip() if model_el.text else (model_el.get('title', '') or '').strip()

                seller_name = ""
                if i < len(sellers):
                    seller_el = sellers[i]
                    seller_name = seller_el.text.strip() if seller_el.text else (seller_el.get('title', '') or '').strip()

                price_raw = ""
                if i < len(prices):
                    price_el = prices[i]
                    price_raw = price_el.text.strip() if price_el.text else ''
                    if not price_raw:
                        price_raw = price_el.text_content().strip() if hasattr(price_el, 'text_content') else ''

                link = ""
                if i < len(links):
                    link_el = links[i]
                    if isinstance(link_el, str):
                        link = link_el
                    else:
                        link = link_el.get('href', '') if hasattr(link_el, 'get') else ''
                        if link and not link.startswith('http'):
                            link = f"https://vako.market{link}"

                cleaned_price = ''.join(c for c in price_raw if c.isdigit() or c in ',.')
                cleaned_price = cleaned_price.replace(',', '.')

                if cleaned_price:
                    results.append({
                        "model": model_name if model_name else "iPhone",
                        "seller": seller_name if seller_name else "Неизвестный магазин",
                        "price": int(float(cleaned_price)),
                        "link": link if link else "—"
                    })

            return results if results else [{"error": "Не найдено"}]

        except requests.exceptions.ConnectionError:
            if attempt < max_retries - 1:
                continue
            return [{"error": "Ошибка соединения"}]

        except requests.exceptions.Timeout:
            if attempt < max_retries - 1:
                continue
            return [{"error": "Таймаут"}]

        except Exception as e:
            return [{"error": f"Ошибка: {str(e)[:50]}"}]

    return [{"error": "Все попытки исчерпаны"}]

# ==========================================================
# 6. ОСНОВНАЯ ЛОГИКА
# ==========================================================

def main():
    print("Запуск мониторинга iPhone через Vako.Market...")
    print("=" * 50)
    start_time = time.time()
    results = []

    for vako_product in vako_products:
        print(f"\n{vako_product['name']}")
        print(f"   URL: {vako_product['url']}")

        vako_data = parse_vako_prices(
            vako_product['url'],
            vako_product['xpath_model'],
            vako_product['xpath_price'],
            vako_product['xpath_seller'],
            vako_product['xpath_link']
        )

        for item in vako_data:
            if "error" in item:
                print(f"  {item['error']}")
            else:
                color = extract_color(item['model'])
                priority = get_color_priority(vako_product['name'], color)
                print(f"  {item['seller']}: {item['price']} руб. ({color}, приоритет: {priority})")
                results.append({
                    "Товар": vako_product['name'],
                    "Цвет": color,
                    "Приоритет цвета": priority,
                    "Магазин": item['seller'],
                    "Цена (руб)": item['price'],
                    "Ссылка": item['link'],
                    "Дата": datetime.now().strftime("%Y-%m-%d %H:%M")
                })

    # --- Сохранение в Excel ---
    filename = f"vako_prices_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Цены"

    headers = ["Товар", "Цвет", "Приоритет цвета", "Магазин", "Цена (руб)", "Ссылка", "Дата"]
    sheet.append(headers)

    for row in results:
        link_display = "Перейти в магазин" if row["Ссылка"] and row["Ссылка"] != "—" else "—"

        sheet.append([
            row["Товар"],
            row["Цвет"],
            row["Приоритет цвета"],
            row["Магазин"],
            row["Цена (руб)"],
            link_display,
            row["Дата"]
        ])

        if row["Ссылка"] and row["Ссылка"] != "—":
            link_cell = sheet.cell(row=sheet.max_row, column=6)
            link_cell.hyperlink = row["Ссылка"]
            link_cell.font = Font(color="0563C1", underline="single")

    # Автоширина колонок
    for col in sheet.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = max_length + 2
        sheet.column_dimensions[column].width = min(adjusted_width, 40)

    wb.save(filename)

    elapsed = time.time() - start_time
    print("\n" + "=" * 50)
    print(f"Готово! Файл сохранён: {filename}")
    print(f"Время выполнения: {elapsed:.2f} сек.")
    print(f"Собрано записей: {len(results)}")

if __name__ == "__main__":
    main()